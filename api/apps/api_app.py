#
#  Copyright 2024 The InfiniFlow Authors. All Rights Reserved.
#
#  Licensed under the Apache License, Version 2.0 (the "License");
#  you may not use this file except in compliance with the License.
#  You may obtain a copy of the License at
#
#      http://www.apache.org/licenses/LICENSE-2.0
#
#  Unless required by applicable law or agreed to in writing, software
#  distributed under the License is distributed on an "AS IS" BASIS,
#  WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
#  See the License for the specific language governing permissions and
#  limitations under the License.
#
from datetime import datetime, timedelta
from quart import request
from api.db.db_models import APIToken, Dialog, Conversation, User,SyncPerson
from api.db.services.api_service import APITokenService, API4ConversationService
from api.db.services.user_service import UserTenantService
from api.utils.api_utils import generate_confirmation_token, get_data_error_result, get_json_result, get_request_json, server_error_response, validate_request
from common.time_utils import current_timestamp, datetime_format
from api.apps import login_required, current_user

import csv
import io
import openpyxl
from flask import Response
# ... 其他已有的 import

@manager.route('/user_dialogs_export_excel_all', methods=['POST'])
async def export_user_dialogs_excel_all():
    import io
    import json
    import time
    from datetime import datetime

    from peewee import JOIN, fn
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
    from openpyxl.utils import get_column_letter

    start_time = time.time()

    # 包括 2026-07-23 当天
    start_date = datetime(2026, 7, 23, 0, 0, 0)

    def parse_json_if_str(value):
        if isinstance(value, str):
            try:
                return json.loads(value)
            except Exception:
                return value

        return value

    def parse_messages(raw_message):
        try:
            message_data = parse_json_if_str(raw_message)

            if isinstance(message_data, list):
                return message_data

            if isinstance(message_data, dict):
                for key in ["messages", "conversation", "history"]:
                    if isinstance(message_data.get(key), list):
                        return message_data[key]

                return [message_data]

            if message_data is None:
                return []

            return [
                {
                    "role": "unknown",
                    "content": str(message_data),
                }
            ]

        except Exception:
            return [
                {
                    "role": "unknown",
                    "content": str(raw_message),
                }
            ]

    def get_dialog_prologue(prompt_config):
        prompt_config = parse_json_if_str(prompt_config)

        if not isinstance(prompt_config, dict):
            return ""

        prologue = prompt_config.get("prologue") or ""

        if isinstance(prologue, (dict, list)):
            return json.dumps(prologue, ensure_ascii=False)

        return str(prologue).strip()

    def get_content(message_item):
        if not isinstance(message_item, dict):
            return str(message_item)

        content = message_item.get("content")

        if content is None:
            content = message_item.get("text")

        if content is None:
            content = message_item.get("message")

        if isinstance(content, (dict, list)):
            return json.dumps(content, ensure_ascii=False)

        if content is None:
            return ""

        return str(content)

    def get_role(message_item):
        role_map = {
            "assistant": "助手",
            "user": "用户",
            "system": "系统",
            "unknown": "未知",
        }

        if not isinstance(message_item, dict):
            return "未知"

        raw_role = str(
            message_item.get("role", "unknown")
        ).strip().lower()

        return role_map.get(raw_role, raw_role)

    def format_datetime(value):
        if not value:
            return ""

        if isinstance(value, datetime):
            return value.strftime("%Y-%m-%d %H:%M:%S")

        return str(value)

    def normalize_sort_datetime(value):
        if not value:
            return datetime.max

        if isinstance(value, datetime):
            return value

        if isinstance(value, str):
            value = value[:19]

            for date_format in [
                "%Y-%m-%d %H:%M:%S",
                "%Y-%m-%d",
                "%Y/%m/%d %H:%M:%S",
                "%Y/%m/%d",
            ]:
                try:
                    return datetime.strptime(value, date_format)
                except Exception:
                    continue

        return datetime.max

    def get_sync_person_value(person, field_name, default=""):
        value = getattr(person, field_name, None)

        if value is None:
            return default

        value = str(value).strip()

        return value or default

    def merge_range(ws, start_row, end_row, column_index, alignment, border):
        if start_row > end_row:
            return

        if start_row < end_row:
            ws.merge_cells(
                start_row=start_row,
                start_column=column_index,
                end_row=end_row,
                end_column=column_index,
            )

        cell = ws.cell(
            row=start_row,
            column=column_index,
        )
        cell.alignment = alignment
        cell.border = border

    def make_safe_filename(value):
        value = str(value or "all")
        return "".join(
            "_" if char in '\\/:*?"<>|' else char
            for char in value
        )

    try:
        # ==========================================================
        # 1. 获取前端参数
        # ==========================================================
        req = await get_request_json()

        if not isinstance(req, dict):
            req = {}

        organization_code = (
            req.get("organizationCode")
            or req.get("organization_code")
            or req.get("department")
            or "全部"
        )

        organization_code = str(organization_code).strip()

        export_all = organization_code in [
            "",
            "全部",
            "all",
            "ALL",
        ]

        allowed_user_ids = None

        # 手机号 -> 部门名称
        phone_department_map = {}

        # 手机号 -> 部门编码
        phone_department_code_map = {}

        # ==========================================================
        # 2. 按部门筛选用户
        # ==========================================================
        if not export_all:
            sync_persons = list(
                SyncPerson
                .select(
                    SyncPerson.phone,
                    SyncPerson.organize,
                    SyncPerson.organizationCode,
                )
                .where(
                    SyncPerson.organizationCode == organization_code
                )
            )

            if not sync_persons:
                return get_data_error_result(
                    message=f"部门 {organization_code} 未找到人员"
                )

            phones = set()

            for person in sync_persons:
                phone = get_sync_person_value(
                    person,
                    "phone",
                )

                if not phone:
                    continue

                phones.add(phone)

                department_name = get_sync_person_value(
                    person,
                    "organize",
                    default=organization_code,
                )

                department_code = get_sync_person_value(
                    person,
                    "organizationCode",
                    default=organization_code,
                )

                phone_department_map[phone] = department_name
                phone_department_code_map[phone] = department_code

            phones = list(phones)

            if not phones:
                return get_data_error_result(
                    message=f"部门 {organization_code} 未找到人员手机号"
                )

            # SyncPerson.phone 对应 User.email
            dept_users = list(
                User
                .select(
                    User.id,
                    User.nickname,
                    User.email,
                )
                .where(
                    fn.TRIM(User.email).in_(phones)
                )
            )

            if not dept_users:
                return get_data_error_result(
                    message=f"部门 {organization_code} 的手机号未匹配到系统用户"
                )

            allowed_user_ids = list(
                {
                    user.id
                    for user in dept_users
                    if user.id
                }
            )

            if not allowed_user_ids:
                return get_data_error_result(
                    message=f"部门 {organization_code} 未匹配到有效用户"
                )

        # ==========================================================
        # 3. 查询用户信息
        # ==========================================================
        user_query = User.select(
            User.id,
            User.nickname,
            User.email,
        )

        if allowed_user_ids is not None:
            user_query = user_query.where(
                User.id.in_(allowed_user_ids)
            )

        user_list = list(user_query)

        users = {
            user.id: {
                "nickname": str(user.nickname or "").strip(),
                "phone": str(user.email or "").strip(),
            }
            for user in user_list
        }

        # ==========================================================
        # 4. 全部导出时，根据手机号反查部门
        # ==========================================================
        if export_all:
            all_user_phones = list(
                {
                    user_info["phone"]
                    for user_info in users.values()
                    if user_info.get("phone")
                }
            )

            if all_user_phones:
                sync_person_query = (
                    SyncPerson
                    .select(
                        SyncPerson.phone,
                        SyncPerson.organize,
                        SyncPerson.organizationCode,
                    )
                    .where(
                        fn.TRIM(SyncPerson.phone).in_(
                            all_user_phones
                        )
                    )
                )

                for person in sync_person_query:
                    phone = get_sync_person_value(
                        person,
                        "phone",
                    )

                    if not phone:
                        continue

                    department_name = get_sync_person_value(
                        person,
                        "organize",
                        default="未分配部门",
                    )

                    department_code = get_sync_person_value(
                        person,
                        "organizationCode",
                    )

                    phone_department_map[phone] = department_name
                    phone_department_code_map[phone] = department_code

        # 给用户补充部门信息
        for user_info in users.values():
            phone = user_info.get("phone") or ""

            user_info["department"] = (
                phone_department_map.get(phone)
                or "未分配部门"
            )

            user_info["department_code"] = (
                phone_department_code_map.get(phone)
                or ""
            )

        # ==========================================================
        # 5. 查询 Conversation 和 Dialog
        # ==========================================================
        conversations_query = (
            Conversation
            .select(
                Conversation.id,
                Conversation.name,
                Conversation.message,
                Conversation.user_id,
                Conversation.dialog_id,
                Conversation.create_date,
                Dialog.id.alias("joined_dialog_id"),
                Dialog.name.alias("dialog_name"),
                Dialog.create_date.alias("dialog_create_date"),
                Dialog.prompt_config.alias("dialog_prompt_config"),
            )
            .join(
                Dialog,
                JOIN.LEFT_OUTER,
                on=(Conversation.dialog_id == Dialog.id),
            )
            .where(
                Conversation.create_date >= start_date
            )
        )

        if allowed_user_ids is not None:
            conversations_query = conversations_query.where(
                Conversation.user_id.in_(allowed_user_ids)
            )

        conversations = list(
            conversations_query
            .dicts()
        )

        # ==========================================================
        # 6. 增加排序字段
        # ==========================================================
        enriched_conversations = []

        for conversation in conversations:
            user_id = conversation.get("user_id") or ""
            user_info = users.get(user_id, {})

            conversation["_user_name"] = (
                user_info.get("nickname")
                or user_id
            )

            conversation["_phone"] = (
                user_info.get("phone")
                or ""
            )

            conversation["_department"] = (
                user_info.get("department")
                or "未分配部门"
            )

            conversation["_department_code"] = (
                user_info.get("department_code")
                or ""
            )

            conversation["_dialog_create_date_sort"] = (
                normalize_sort_datetime(
                    conversation.get("dialog_create_date")
                )
            )

            conversation["_conversation_create_date_sort"] = (
                normalize_sort_datetime(
                    conversation.get("create_date")
                )
            )

            enriched_conversations.append(conversation)

        # 排序：
        # 部门 -> 用户 -> 主题 -> Conversation
        enriched_conversations.sort(
            key=lambda item: (
                item.get("_department") or "",
                item.get("_department_code") or "",
                item.get("_user_name") or "",
                item.get("user_id") or "",
                item.get("_dialog_create_date_sort"),
                item.get("dialog_name") or "",
                item.get("dialog_id") or "",
                item.get("_conversation_create_date_sort"),
                item.get("name") or "",
                item.get("id") or "",
            )
        )

        # ==========================================================
        # 7. 创建工作簿
        # ==========================================================
        wb = Workbook()
        ws = wb.active
        ws.title = "对话记录"

        headers = [
            "部门",
            "用户名称",
            "手机号码",
            "主题名称",
            "主题创建时间",
            "对话名称",
            "对话创建时间",
            "消息序号",
            "发送者",
            "消息详情",
        ]

        ws.append(headers)

        header_fill = PatternFill(
            "solid",
            fgColor="D9EAF7",
        )

        header_font = Font(
            bold=True,
        )

        center_alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )

        left_alignment = Alignment(
            horizontal="left",
            vertical="top",
            wrap_text=True,
        )

        thin_side = Side(
            style="thin",
            color="CCCCCC",
        )

        border = Border(
            left=thin_side,
            right=thin_side,
            top=thin_side,
            bottom=thin_side,
        )

        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = center_alignment
            cell.border = border

        current_row = 2
        data_start_row = current_row

        conversation_count = 0
        message_count = 0

        # 保存每个层级的行范围
        department_ranges = {}
        user_ranges = {}
        dialog_ranges = {}

        # ==========================================================
        # 8. 写入消息
        # ==========================================================
        for conversation in enriched_conversations:
            conversation_count += 1
            conversation_start_row = current_row

            user_id = conversation.get("user_id") or ""

            user_name = (
                conversation.get("_user_name")
                or user_id
                or "未知用户"
            )

            user_phone = conversation.get("_phone") or ""

            department = (
                conversation.get("_department")
                or "未分配部门"
            )

            dialog_id = (
                conversation.get("dialog_id")
                or conversation.get("joined_dialog_id")
                or ""
            )

            dialog_name = (
                conversation.get("dialog_name")
                or "未命名主题"
            )

            dialog_create_date = format_datetime(
                conversation.get("dialog_create_date")
            )

            conversation_name = (
                conversation.get("name")
                or "新对话"
            )

            conversation_create_date = format_datetime(
                conversation.get("create_date")
            )

            messages = parse_messages(
                conversation.get("message")
            )

            # 补第一句助手消息
            prologue = get_dialog_prologue(
                conversation.get("dialog_prompt_config")
            )

            if prologue:
                has_same_first_assistant = False

                if messages:
                    first_message = messages[0]

                    if isinstance(first_message, dict):
                        first_role = str(
                            first_message.get("role", "")
                        ).strip().lower()

                        first_content = get_content(
                            first_message
                        ).strip()

                        has_same_first_assistant = (
                            first_role == "assistant"
                            and first_content == prologue
                        )

                if not has_same_first_assistant:
                    messages.insert(
                        0,
                        {
                            "role": "assistant",
                            "content": prologue,
                        },
                    )

            if not messages:
                messages = [
                    {
                        "role": "",
                        "content": "",
                    }
                ]

            for index, message_item in enumerate(
                messages,
                start=1,
            ):
                role = get_role(message_item)

                content = get_content(
                    message_item
                ).replace("\r", "")

                ws.append(
                    [
                        department,
                        user_name,
                        user_phone,
                        dialog_name,
                        dialog_create_date,
                        conversation_name,
                        conversation_create_date,
                        index,
                        role,
                        content,
                    ]
                )

                for column_index in range(1, 11):
                    cell = ws.cell(
                        row=current_row,
                        column=column_index,
                    )

                    cell.border = border

                    # 第 10 列是消息详情
                    cell.alignment = (
                        left_alignment
                        if column_index == 10
                        else center_alignment
                    )

                current_row += 1
                message_count += 1

            conversation_end_row = current_row - 1

            # ======================================================
            # Conversation 维度合并
            # 对话名称：第 6 列
            # 对话创建时间：第 7 列
            # ======================================================
            for column_index in [6, 7]:
                merge_range(
                    ws,
                    conversation_start_row,
                    conversation_end_row,
                    column_index,
                    center_alignment,
                    border,
                )

            # ======================================================
            # 部门维度记录
            # 相同部门的所有记录连续合并
            # ======================================================
            department_key = department or "未分配部门"

            if department_key not in department_ranges:
                department_ranges[department_key] = [
                    conversation_start_row,
                    conversation_end_row,
                ]
            else:
                department_ranges[department_key][1] = conversation_end_row

            # ======================================================
            # 用户维度记录
            # 同一个部门下的同一个用户合并
            # ======================================================
            user_key = (
                f"{department_key}::{user_id or user_name}"
            )

            if user_key not in user_ranges:
                user_ranges[user_key] = [
                    conversation_start_row,
                    conversation_end_row,
                ]
            else:
                user_ranges[user_key][1] = conversation_end_row

            # ======================================================
            # 主题维度记录
            # 同一个用户下的同一个主题合并
            # ======================================================
            dialog_key = (
                f"{user_key}::{dialog_id or dialog_name}"
            )

            if dialog_key not in dialog_ranges:
                dialog_ranges[dialog_key] = [
                    conversation_start_row,
                    conversation_end_row,
                ]
            else:
                dialog_ranges[dialog_key][1] = conversation_end_row

        data_end_row = current_row - 1

        # ==========================================================
        # 9. 合并部门
        # 部门是第 1 列
        # ==========================================================
        for start_row, end_row in department_ranges.values():
            merge_range(
                ws,
                start_row,
                end_row,
                1,
                center_alignment,
                border,
            )

        # ==========================================================
        # 10. 合并用户
        # 用户名称是第 2 列
        # 手机号码是第 3 列
        # ==========================================================
        for start_row, end_row in user_ranges.values():
            for column_index in [2, 3]:
                merge_range(
                    ws,
                    start_row,
                    end_row,
                    column_index,
                    center_alignment,
                    border,
                )

        # ==========================================================
        # 11. 合并主题
        # 主题名称是第 4 列
        # 主题创建时间是第 5 列
        # ==========================================================
        for start_row, end_row in dialog_ranges.values():
            for column_index in [4, 5]:
                merge_range(
                    ws,
                    start_row,
                    end_row,
                    column_index,
                    center_alignment,
                    border,
                )

        # ==========================================================
        # 12. 设置列宽
        # ==========================================================
        column_widths = {
            1: 28,   # 部门
            2: 20,   # 用户名称
            3: 18,   # 手机号码
            4: 28,   # 主题名称
            5: 20,   # 主题创建时间
            6: 40,   # 对话名称
            7: 20,   # 对话创建时间
            8: 10,   # 消息序号
            9: 15,   # 发送者
            10: 90,  # 消息详情
        }

        for column_index, width in column_widths.items():
            ws.column_dimensions[
                get_column_letter(column_index)
            ].width = width

        ws.freeze_panes = "A2"

        # ==========================================================
        # 13. 导出 Excel
        # ==========================================================
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)

        export_scope = (
            "all"
            if export_all
            else organization_code
        )

        filename = (
            f"dialog_logs_{export_scope}_from_20260723_"
            f"{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        )

        response = Response(
            output.getvalue(),
            mimetype=(
                "application/vnd.openxmlformats-officedocument."
                "spreadsheetml.sheet"
            ),
            headers={
                "Content-Disposition": (
                    f"attachment; filename={filename}"
                )
            },
        )

        print(f"导出完成: {filename}")
        print(
            f"导出范围: "
            f"{'全部' if export_all else organization_code}"
        )
        print(f"Conversation 数量: {conversation_count}")
        print(f"消息数量: {message_count}")
        print(f"后端耗时: {time.time() - start_time:.2f} 秒")

        output.close()

        return response

    except Exception as e:
        print(f"Export Excel Error: {e}")
        return server_error_response(e)
    
@manager.route('/user_dialogs_export_excel', methods=['POST'])
async def export_user_dialogs_excel():
    import io
    import json
    import time
    from datetime import datetime

    from peewee import fn
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
    from openpyxl.utils import get_column_letter

    start_time = time.time()

    def parse_json_if_str(value):
        if isinstance(value, str):
            try:
                return json.loads(value)
            except Exception:
                return value

        return value

    def parse_messages(raw_message):
        """
        兼容以下格式：

        [
            {"role": "user", "content": "..."},
            {"role": "assistant", "content": "..."}
        ]

        或：

        {
            "messages": [...]
        }
        """
        try:
            msg_data = parse_json_if_str(raw_message)

            if isinstance(msg_data, list):
                return msg_data

            if isinstance(msg_data, dict):
                for key in ["messages", "conversation", "history"]:
                    if isinstance(msg_data.get(key), list):
                        return msg_data[key]

                # 如果本身就是一条消息
                return [msg_data]

            if msg_data is None:
                return []

            return [
                {
                    "role": "unknown",
                    "content": str(msg_data),
                }
            ]

        except Exception:
            return [
                {
                    "role": "unknown",
                    "content": str(raw_message),
                }
            ]

    def get_dialog_prologue(prompt_config):
        """
        获取主题配置中的开场白。
        通常第一句助手消息存储在：
        Dialog.prompt_config["prologue"]
        """
        prompt_config = parse_json_if_str(prompt_config)

        if not isinstance(prompt_config, dict):
            return ""

        prologue = prompt_config.get("prologue") or ""

        if isinstance(prologue, (dict, list)):
            return json.dumps(prologue, ensure_ascii=False)

        return str(prologue).strip()

    def get_content(msg):
        if not isinstance(msg, dict):
            return str(msg)

        content = msg.get("content")

        if content is None:
            content = msg.get("text")

        if content is None:
            content = msg.get("message")

        if isinstance(content, (dict, list)):
            return json.dumps(content, ensure_ascii=False)

        if content is None:
            return ""

        return str(content)

    def get_role(msg):
        role_map = {
            "assistant": "助手",
            "user": "用户",
            "system": "系统",
            "unknown": "未知",
        }

        if not isinstance(msg, dict):
            return "未知"

        raw_role = str(
            msg.get("role", "unknown")
        ).strip().lower()

        return role_map.get(raw_role, raw_role)

    def format_datetime(value):
        if not value:
            return ""

        if isinstance(value, datetime):
            return value.strftime("%Y-%m-%d %H:%M:%S")

        return str(value)

    def merge_range(ws, start_row, end_row, column_index, alignment, border):
        if start_row > end_row:
            return

        if start_row < end_row:
            ws.merge_cells(
                start_row=start_row,
                start_column=column_index,
                end_row=end_row,
                end_column=column_index,
            )

        cell = ws.cell(
            row=start_row,
            column=column_index,
        )
        cell.alignment = alignment
        cell.border = border

    try:
        # ==========================================================
        # 1. 获取请求参数
        # ==========================================================
        req = await get_request_json()

        if not isinstance(req, dict):
            req = {}

        tenant_id = req.get("tenant_id")
        dialog_id = req.get("dialog_id")

        if not tenant_id:
            return get_data_error_result(
                message="Tenant_id not found!"
            )

        if not dialog_id:
            return get_data_error_result(
                message="Dialog_id not found!"
            )

        # ==========================================================
        # 2. 查询用户
        # ==========================================================
        user = (
            User
            .select(
                User.id,
                User.nickname,
                User.email,
            )
            .where(User.id == tenant_id)
            .first()
        )

        if not user:
            return get_data_error_result(
                message="User not found!"
            )

        user_name = str(
            user.nickname or user.id
        ).strip()

        # 当前业务中 User.email 保存的是手机号
        user_phone = str(
            user.email or ""
        ).strip()

        # ==========================================================
        # 3. 根据手机号查询部门
        #
        # User.email
        #     -> SyncPerson.phone
        #
        # SyncPerson.organize
        #     -> 部门名称
        #
        # SyncPerson.organizationCode
        #     -> 部门编码
        # ==========================================================
        department = ""
        department_code = ""

        if user_phone:
            sync_person = (
                SyncPerson
                .select(
                    SyncPerson.phone,
                    SyncPerson.organize,
                    SyncPerson.organizationCode,
                )
                .where(
                    fn.TRIM(SyncPerson.phone) == user_phone
                )
                .first()
            )

            if sync_person:
                department = str(
                    sync_person.organize or ""
                ).strip()

                department_code = str(
                    sync_person.organizationCode or ""
                ).strip()

        # 如果没有查到部门，显示未分配部门
        department_display = department or "未分配部门"

        # 临时调试日志，可以确认手机号和部门是否匹配
        print("User ID:", repr(user.id))
        print("User email/phone:", repr(user_phone))
        print("Department:", repr(department))
        print("Department code:", repr(department_code))

        # ==========================================================
        # 4. 查询主题
        # ==========================================================
        dialog = (
            Dialog
            .select(
                Dialog.id,
                Dialog.name,
                Dialog.tenant_id,
                Dialog.create_date,
                Dialog.prompt_config,
            )
            .where(
                (Dialog.id == dialog_id)
                & (Dialog.tenant_id == tenant_id)
            )
            .first()
        )

        if not dialog:
            return get_data_error_result(
                message="Dialog not found!"
            )

        dialog_name = str(
            dialog.name or ""
        ).strip()

        dialog_create_date = format_datetime(
            dialog.create_date
        )

        dialog_prologue = get_dialog_prologue(
            dialog.prompt_config
        )

        # ==========================================================
        # 5. 查询 Conversation
        # ==========================================================
        conversations = list(
            Conversation
            .select(
                Conversation.id,
                Conversation.name,
                Conversation.message,
                Conversation.create_date,
            )
            .where(
                (Conversation.dialog_id == dialog_id)
                & (Conversation.user_id == tenant_id)
            )
            .order_by(
                Conversation.create_date.asc(),
                Conversation.id.asc(),
            )
        )

        # ==========================================================
        # 6. 创建 Excel
        # ==========================================================
        wb = Workbook()
        ws = wb.active
        ws.title = "对话记录"

        # 不导出用户 ID、主题 ID、Conversation ID
        headers = [
            "部门",
            "用户名称",
            
            "手机号码",
            "主题名称",
            "主题创建时间",
            "对话名称",
            "对话创建时间",
            "消息序号",
            "发送者",
            "消息详情",
        ]

        ws.append(headers)

        # ==========================================================
        # 7. 设置 Excel 样式
        # ==========================================================
        header_fill = PatternFill(
            "solid",
            fgColor="D9EAF7",
        )

        header_font = Font(
            bold=True,
        )

        center_alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )

        left_alignment = Alignment(
            horizontal="left",
            vertical="top",
            wrap_text=True,
        )

        thin_side = Side(
            style="thin",
            color="CCCCCC",
        )

        border = Border(
            left=thin_side,
            right=thin_side,
            top=thin_side,
            bottom=thin_side,
        )

        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = center_alignment
            cell.border = border

        current_row = 2
        data_start_row = current_row
        conversation_count = 0
        message_count = 0

        # ==========================================================
        # 8. 写入对话数据
        # ==========================================================
        for conversation in conversations:
            conversation_count += 1

            conversation_start_row = current_row

            conversation_name = str(
                conversation.name or ""
            ).strip()

            conversation_create_date = format_datetime(
                conversation.create_date
            )

            messages = parse_messages(
                conversation.message
            )

            # 补第一句助手消息
            # 如果 Conversation.message 中已经有相同的第一句，
            # 则不重复添加。
            if dialog_prologue:
                has_same_first_assistant = False

                if messages:
                    first_message = messages[0]

                    if isinstance(first_message, dict):
                        first_role = str(
                            first_message.get("role", "")
                        ).strip().lower()

                        first_content = get_content(
                            first_message
                        ).strip()

                        has_same_first_assistant = (
                            first_role == "assistant"
                            and first_content == dialog_prologue
                        )

                if not has_same_first_assistant:
                    messages.insert(
                        0,
                        {
                            "role": "assistant",
                            "content": dialog_prologue,
                        },
                    )

            if not messages:
                messages = [
                    {
                        "role": "",
                        "content": "",
                    }
                ]

            for index, message_item in enumerate(
                messages,
                start=1,
            ):
                role = get_role(message_item)

                content = get_content(
                    message_item
                ).replace("\r", "")

                ws.append(
                    [   
                        department_display,
                        user_name,
                        
                        user_phone,
                        dialog_name,
                        dialog_create_date,
                        conversation_name,
                        conversation_create_date,
                        index,
                        role,
                        content,
                    ]
                )

                # 当前 Excel 一共 10 列
                for column_index in range(1, 11):
                    cell = ws.cell(
                        row=current_row,
                        column=column_index,
                    )

                    cell.border = border

                    # 第 10 列是消息详情
                    cell.alignment = (
                        left_alignment
                        if column_index == 10
                        else center_alignment
                    )

                current_row += 1
                message_count += 1

            conversation_end_row = current_row - 1

            # 同一个 Conversation 合并：
            # 第 6 列：对话名称
            # 第 7 列：对话创建时间
            for column_index in [6, 7]:
                merge_range(
                    ws,
                    conversation_start_row,
                    conversation_end_row,
                    column_index,
                    center_alignment,
                    border,
                )

        data_end_row = current_row - 1

        # ==========================================================
        # 9. 合并用户和主题信息
        # ==========================================================
        if data_start_row <= data_end_row:
            # 当前是单个用户、单个主题导出
            #
            # 第 1 列：部门
            # 第 2 列：用户名称
            # 第 3 列：手机号码
            for column_index in [1, 2, 3]:
                merge_range(
                    ws,
                    data_start_row,
                    data_end_row,
                    column_index,
                    center_alignment,
                    border,
                )

            # 第 4 列：主题名称
            # 第 5 列：主题创建时间
            for column_index in [4, 5]:
                merge_range(
                    ws,
                    data_start_row,
                    data_end_row,
                    column_index,
                    center_alignment,
                    border,
                )

        # ==========================================================
        # 10. 设置列宽
        # ==========================================================
        column_widths = {
            1: 28,   # 部门
            2: 20,   # 用户名称
            3: 18,   # 手机号码
            4: 28,   # 主题名称
            5: 20,   # 主题创建时间
            6: 40,   # 对话名称
            7: 20,   # 对话创建时间
            8: 10,   # 消息序号
            9: 15,   # 发送者
            10: 90,  # 消息详情
        }

        for column_index, width in column_widths.items():
            ws.column_dimensions[
                get_column_letter(column_index)
            ].width = width

        ws.freeze_panes = "A2"

        # ==========================================================
        # 11. 输出 Excel
        # ==========================================================
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)

        filename = (
            f"dialog_logs_{dialog_id}_"
            f"{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        )

        response = Response(
            output.getvalue(),
            mimetype=(
                "application/vnd.openxmlformats-officedocument."
                "spreadsheetml.sheet"
            ),
            headers={
                "Content-Disposition": (
                    f"attachment; filename={filename}"
                )
            },
        )

        print(f"导出完成: {filename}")
        print(f"用户: {user_name}")
        print(f"部门: {department_display}")
        print(f"部门编码: {department_code}")
        print(f"手机号: {user_phone}")
        print(f"主题: {dialog_name}")
        print(f"Conversation 数量: {conversation_count}")
        print(f"消息数量: {message_count}")
        print(
            f"后端耗时: "
            f"{time.time() - start_time:.2f} 秒"
        )

        output.close()

        return response

    except Exception as e:
        print(f"Export Excel Error: {e}")
        return server_error_response(e)

@manager.route('/user_dialogs_and_conversations', methods=['POST'])
# @login_required
async def get_user_dialogs_and_conversations():
    from datetime import datetime

    req = await get_request_json()
    tenant_id = req.get("tenant_id")

    if not tenant_id:
        return get_data_error_result(message="Tenant_id not found!")

    def format_datetime(value):
        if not value:
            return ""

        if isinstance(value, datetime):
            return value.strftime("%Y-%m-%d %H:%M:%S")

        return str(value)

    try:
        # 1. 查询该租户下的所有 Dialog，按照主题创建时间排序
        dialogs = (
            Dialog
            .select(
                Dialog.id,
                Dialog.name,
                Dialog.create_date
            )
            .where(Dialog.tenant_id == tenant_id)
            .order_by(
                Dialog.create_date.asc(),
                Dialog.id.asc()
            )
        )

        # 2. 查询该租户下的所有 Conversation，按照对话创建时间排序
        conversations = (
            Conversation
            .select(
                Conversation.id,
                Conversation.name,
                Conversation.message,
                Conversation.dialog_id,
                Conversation.create_date
            )
            .where(Conversation.user_id == tenant_id)
            .order_by(
                Conversation.create_date.asc(),
                Conversation.id.asc()
            )
        )

        # 3. 将 Conversation 转换为以 dialog_id 为键的字典
        # 因为上面已经按照 create_date 排序，所以 append 后每个主题里的对话也是有序的
        conv_map = {}

        for c in conversations:
            if c.dialog_id not in conv_map:
                conv_map[c.dialog_id] = []

            conv_map[c.dialog_id].append({
                "id": c.id,
                "name": c.name,
                "message": c.message,
                "create_date": format_datetime(c.create_date),
                "conversation_create_date": format_datetime(c.create_date),
            })

        # 4. 组装最终的树状结构
        # dialogs 本身已经按主题创建时间排序
        dialog_list = []

        for d in dialogs:
            dialog_list.append({
                "id": d.id,
                "name": d.name,
                "create_date": format_datetime(d.create_date),
                "dialog_create_date": format_datetime(d.create_date),
                "conversations": conv_map.get(d.id, [])
            })

        return get_json_result(data=dialog_list)

    except Exception as e:
        return server_error_response(e)
    
@manager.route('/group_member_stats', methods=['POST'])
# @login_required
async def group_member_stats():
    """
    获取所有组及组内成员的 Token 和问答次数统计

    POST body:
    {
        "period": "all" | "day" | "week" | "month" | "year"
    }
    """
    try:
        from datetime import datetime, timedelta
        from api.db.db_models import APIToken
        from api.db.db_models import Group, UserGroup, User
        from peewee import fn

        req = await get_request_json()
        req = req or {}

        period = req.get("period", "all")

        if period not in ["all", "day", "week", "month", "year"]:
            period = "all"

        now = datetime.now()

        start_time = None
        end_time = None
        start_time_ts = None
        end_time_ts = None

        if period == "day":
            start_time = now.replace(hour=0, minute=0, second=0, microsecond=0)
            end_time = now

        elif period == "week":
            start_time = now - timedelta(days=now.weekday())
            start_time = start_time.replace(hour=0, minute=0, second=0, microsecond=0)
            end_time = now

        elif period == "month":
            start_time = now.replace(day=1, hour=0, minute=0, second=0, microsecond=0)
            end_time = now

        elif period == "year":
            start_time = now.replace(month=1, day=1, hour=0, minute=0, second=0, microsecond=0)
            end_time = now

        elif period == "all":
            start_time = None
            end_time = None

        if start_time and end_time:
            start_time_ts = int(start_time.timestamp() * 1000)
            end_time_ts = int(end_time.timestamp() * 1000)

        def ms_to_date(ms):
            return datetime.fromtimestamp(int(ms) / 1000).strftime("%Y-%m-%d")

        def ms_to_datetime(ms):
            return datetime.fromtimestamp(int(ms) / 1000)

        def build_day_items_by_range(range_start, range_end):
            if not range_start or not range_end:
                return []

            current_day = range_start.replace(hour=0, minute=0, second=0, microsecond=0)
            last_day = range_end.replace(hour=0, minute=0, second=0, microsecond=0)

            days = []

            while current_day <= last_day:
                days.append({
                    "date": current_day.strftime("%Y-%m-%d"),
                    "tokens": 0
                })
                current_day += timedelta(days=1)

            return days

        def build_week_day_items(week_start):
            days = []

            for i in range(7):
                day = week_start + timedelta(days=i)
                days.append({
                    "date": day.strftime("%Y-%m-%d"),
                    "tokens": 0
                })

            return days

        groups = Group.select()
        final_result = []

        for group in groups:
            members = list(
                UserGroup
                .select(
                    UserGroup.user_id,
                    User.nickname
                )
                .join(User, on=(UserGroup.user_id == User.id))
                .where(UserGroup.group_id == group.group_id)
                .dicts()
            )

            tenant_ids = [m["user_id"] for m in members]

            member_list = []
            total_tokens = 0
            total_dialogs = 0
            daily_tokens = []

            if tenant_ids:
                conditions = [
                    APIToken.tenant_id.in_(tenant_ids)
                ]

                if period != "all":
                    conditions.extend([
                        APIToken.create_time >= start_time_ts,
                        APIToken.create_time <= end_time_ts
                    ])

                member_stats_list = (
                    APIToken
                    .select(
                        APIToken.tenant_id,
                        fn.SUM(APIToken.token).alias("total_tokens"),
                        fn.COUNT(APIToken.dialog_id).alias("dialog_count")
                    )
                    .where(*conditions)
                    .group_by(APIToken.tenant_id)
                    .dicts()
                )

                stats_map = {
                    stat["tenant_id"]: stat
                    for stat in member_stats_list
                }

                group_token_rows = list(
                    APIToken
                    .select(
                        APIToken.create_time,
                        APIToken.token
                    )
                    .where(*conditions)
                    .dicts()
                )

                if period == "week" and start_time:
                    daily_tokens = build_week_day_items(start_time)

                elif period != "all":
                    daily_tokens = build_day_items_by_range(start_time, end_time)

                else:
                    create_times = [
                        int(row["create_time"])
                        for row in group_token_rows
                        if row.get("create_time")
                    ]

                    if create_times:
                        group_start_time = ms_to_datetime(min(create_times))
                        group_end_time = ms_to_datetime(max(create_times))
                        daily_tokens = build_day_items_by_range(
                            group_start_time,
                            group_end_time
                        )
                    else:
                        daily_tokens = []

                daily_token_map = {
                    item["date"]: item
                    for item in daily_tokens
                }

                for row in group_token_rows:
                    create_time = row.get("create_time")
                    token = int(row.get("token") or 0)

                    if not create_time:
                        continue

                    day_key = ms_to_date(create_time)

                    if day_key in daily_token_map:
                        daily_token_map[day_key]["tokens"] += token

                for m in members:
                    user_id = m["user_id"]
                    nickname = m.get("nickname") or ""

                    stat = stats_map.get(
                        user_id,
                        {
                            "total_tokens": 0,
                            "dialog_count": 0
                        }
                    )

                    tokens = int(stat["total_tokens"] or 0)
                    dialogs = int(stat["dialog_count"] or 0)

                    total_tokens += tokens
                    total_dialogs += dialogs

                    member_list.append({
                        "user_id": user_id,
                        "nickname": nickname,
                        "token_usage": tokens,
                        "dialog_count": dialogs
                    })

            final_result.append({
                "group_id": group.group_id,
                "group_name": group.group_name,
                "period": period,

                "start_time": start_time.strftime("%Y-%m-%d %H:%M:%S") if start_time else None,
                "end_time": end_time.strftime("%Y-%m-%d %H:%M:%S") if end_time else None,

                "start_time_ts": start_time_ts,
                "end_time_ts": end_time_ts,

                "total_tokens": total_tokens,
                "total_dialogs": total_dialogs,
                "daily_tokens": daily_tokens,
                "members": member_list
            })

        group_order = {
            "工艺研究一室": 0,
            "工艺研究二室": 1,
            "工艺研究三室": 2,
            "新品事业部研发部": 3,
        }


        final_result.sort(
            key=lambda item: group_order.get(item.get("group_name"), 999)
        )

        print(final_result)
        return get_json_result(data=final_result)

    except Exception as e:
        return server_error_response(e)


@manager.route('/new_token', methods=['POST'])  # noqa: F821
@login_required
async def new_token():
    req = await get_request_json()
    try:
        tenants = UserTenantService.query(user_id=current_user.id)
        if not tenants:
            return get_data_error_result(message="Tenant not found!")

        tenant_id = tenants[0].tenant_id
        obj = {"tenant_id": tenant_id, "token": generate_confirmation_token(),
               "create_time": current_timestamp(),
               "create_date": datetime_format(datetime.now()),
               "update_time": None,
               "update_date": None
               }
        if req.get("canvas_id"):
            obj["dialog_id"] = req["canvas_id"]
            obj["source"] = "agent"
        else:
            obj["dialog_id"] = req["dialog_id"]

        if not APITokenService.save(**obj):
            return get_data_error_result(message="Fail to new a dialog!")

        return get_json_result(data=obj)
    except Exception as e:
        return server_error_response(e)


@manager.route('/token_list', methods=['GET'])  # noqa: F821
@login_required
def token_list():
    try:
        tenants = UserTenantService.query(user_id=current_user.id)
        if not tenants:
            return get_data_error_result(message="Tenant not found!")

        id = request.args["dialog_id"] if "dialog_id" in request.args else request.args["canvas_id"]
        objs = APITokenService.query(tenant_id=tenants[0].tenant_id, dialog_id=id)
        return get_json_result(data=[o.to_dict() for o in objs])
    except Exception as e:
        return server_error_response(e)


@manager.route('/rm', methods=['POST'])  # noqa: F821
@validate_request("tokens", "tenant_id")
@login_required
async def rm():
    req = await get_request_json()
    try:
        for token in req["tokens"]:
            APITokenService.filter_delete(
                [APIToken.tenant_id == req["tenant_id"], APIToken.token == token])
        return get_json_result(data=True)
    except Exception as e:
        return server_error_response(e)


@manager.route('/stats', methods=['GET'])  # noqa: F821
@login_required
def stats():
    try:
        tenants = UserTenantService.query(user_id=current_user.id)
        if not tenants:
            return get_data_error_result(message="Tenant not found!")
        objs = API4ConversationService.stats(
            tenants[0].tenant_id,
            request.args.get(
                "from_date",
                (datetime.now() -
                 timedelta(
                     days=7)).strftime("%Y-%m-%d 00:00:00")),
            request.args.get(
                "to_date",
                datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
            "agent" if "canvas_id" in request.args else None)

        res = {"pv": [], "uv": [], "speed": [], "tokens": [], "round": [], "thumb_up": []}

        for obj in objs:
            dt = obj["dt"]
            res["pv"].append((dt, obj["pv"]))
            res["uv"].append((dt, obj["uv"]))
            res["speed"].append((dt, float(obj["tokens"]) / (float(obj["duration"]) + 0.1))) # +0.1 to avoid division by zero
            res["tokens"].append((dt, float(obj["tokens"]) / 1000.0)) # convert to thousands
            res["round"].append((dt, obj["round"]))
            res["thumb_up"].append((dt, obj["thumb_up"]))

        return get_json_result(data=res)
    except Exception as e:
        return server_error_response(e)


import os
import hashlib
import jwt
import httpx

from quart import request, send_file

ONLYOFFICE_JWT_SECRET = os.getenv(
    "ONLYOFFICE_JWT_SECRET",
    "3b6167155dfdd4ea76ea9bf9852858af29978e9f862215971d17b4a5b3cb6324"
)

# OnlyOffice 容器访问你的后端用这个地址
INTERNAL_API_BASE_URL = "http://host.docker.internal:9380/v1/api"

# 固定测试文件
TEST_DOCX_PATH = "/home/zyb/onlyoffice-test-docs/test.docx"


def get_test_doc_key():
    if not os.path.exists(TEST_DOCX_PATH):
        return "test-docx-not-exists"

    mtime = str(os.path.getmtime(TEST_DOCX_PATH))
    raw = f"test-docx-{mtime}"

    return hashlib.md5(raw.encode("utf-8")).hexdigest()


@manager.route('/onlyoffice/ping', methods=['GET'])  # noqa: F821
async def onlyoffice_ping():
    return {
        "code": 0,
        "message": "success",
        "data": {
            "message": "onlyoffice route ok"
        }
    }


@manager.route('/onlyoffice/editor-config/test', methods=['GET'])  # noqa: F821
async def onlyoffice_editor_config_test():
    """
    前端请求这个接口，获取 OnlyOffice 编辑器配置。

    真实地址：
    GET /v1/api/onlyoffice/editor-config/test
    """
    try:
        if not os.path.exists(TEST_DOCX_PATH):
            return {
                "code": 404,
                "message": f"测试文件不存在: {TEST_DOCX_PATH}",
                "data": None
            }

        config = {
            "documentType": "word",
            "width": "100%",
            "height": "100%",
            "document": {
                "fileType": "docx",
                "key": get_test_doc_key(),
                "title": "test.docx",
                "url": f"{INTERNAL_API_BASE_URL}/onlyoffice/files/test",
                "permissions": {
                    "edit": True,
                    "download": True,
                    "print": True,
                    "review": True,
                    "comment": True
                }
            },
            "editorConfig": {
                "mode": "edit",
                "lang": "zh-CN",
                "callbackUrl": f"{INTERNAL_API_BASE_URL}/onlyoffice/callback/test",
                "user": {
                    "id": "user-001",
                    "name": "测试用户"
                },
                "customization": {
                    "autosave": True,
                    "forcesave": True
                }
            }
        }

        token = jwt.encode(
            config,
            ONLYOFFICE_JWT_SECRET,
            algorithm="HS256"
        )

        config["token"] = token

        return {
            "code": 0,
            "message": "success",
            "data": config
        }

    except Exception as e:
        return {
            "code": 500,
            "message": str(e),
            "data": None
        }


@manager.route('/onlyoffice/files/test', methods=['GET'])  # noqa: F821
async def onlyoffice_get_file_test():
    """
    OnlyOffice 下载 test.docx。

    真实地址：
    GET /v1/api/onlyoffice/files/test

    注意：这个接口不要加 login_required。
    """
    try:
        if not os.path.exists(TEST_DOCX_PATH):
            return {
                "code": 404,
                "message": f"文件不存在: {TEST_DOCX_PATH}",
                "data": None
            }

        return await send_file(
            TEST_DOCX_PATH,
            mimetype="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            as_attachment=False
        )

    except Exception as e:
        return {
            "code": 500,
            "message": str(e),
            "data": None
        }


@manager.route('/onlyoffice/callback/test', methods=['POST'])  # noqa: F821
async def onlyoffice_callback_test():
    """
    OnlyOffice 保存回调。

    真实地址：
    POST /v1/api/onlyoffice/callback/test

    注意：这个接口不要加 login_required。
    OnlyOffice 要求返回 {"error": 0}
    """
    try:
        body = await request.get_json()
        body = body or {}

        print("OnlyOffice callback body:", body)

        status = body.get("status")

        if status in [2, 6]:
            download_url = body.get("url")

            if not download_url:
                print("OnlyOffice callback 没有 url")
                return {"error": 1}

            async with httpx.AsyncClient(timeout=120) as client:
                response = await client.get(download_url)
                response.raise_for_status()

            tmp_path = TEST_DOCX_PATH + ".tmp"

            with open(tmp_path, "wb") as f:
                f.write(response.content)

            os.replace(tmp_path, TEST_DOCX_PATH)

            print("OnlyOffice 文档已保存:", TEST_DOCX_PATH)

        return {"error": 0}

    except Exception as e:
        print("OnlyOffice callback error:", e)
        return {"error": 1}


@manager.route('/onlyoffice/download/test', methods=['GET'])  # noqa: F821
async def onlyoffice_download_test():
    """
    浏览器下载编辑后的 DOCX。

    真实地址：
    GET /v1/api/onlyoffice/download/test
    """
    return await send_file(
        TEST_DOCX_PATH,
        mimetype="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        as_attachment=True
    )

